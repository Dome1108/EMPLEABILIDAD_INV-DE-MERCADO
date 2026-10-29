from pathlib import Path
import math
import openpyxl
import xlsxwriter

root=Path(__file__).parent
source=openpyxl.load_workbook(root.parent/'Tabla de valores tendencias.xlsx',data_only=True)
path=root/'Modelo_comparativo_viabilidad_v4.xlsx'
w=xlsxwriter.Workbook(path)
head=w.add_format({'bold':True,'bg_color':'#006DB5','font_color':'white','text_wrap':True})
number=w.add_format({'num_format':'0.00"%"'})
percent=w.add_format({'num_format':'0.0%'})
wrap=w.add_format({'text_wrap':True,'valign':'top'})
for sheet in source:
    dest=w.add_worksheet('Fuente_'+sheet.title)
    for i,row in enumerate(sheet.values):
        for j,v in enumerate(row):
            if v is not None:dest.write(i,j,v)
    dest.set_column(0,19,20)
p=w.add_worksheet('Parámetros')
params=[('Demanda',.30),('Empleabilidad',.22),('Oportunidad',.28),('Precio',.12),('Sector',.08),('Nombre dentro de demanda',.65),('Búsquedas dentro del nombre',.90),('Matriculados dentro de demanda educativa',.70),('Vacantes/profesionales dentro de empleabilidad',.40),('Alineación dentro de empleabilidad',.35),('Límite inferior precio/media',.85),('Límite superior precio/media',1.10),('Penalización precio: puntos por unidad de desviación',500),('Umbral comparable',90),('Umbral intermedio',75),('Umbral inferior',60),('Umbral búsquedas para cuadrante verde (índice)',75),('Máximo competidores para cuadrante verde',10)]
p.write_row(0,0,['Parámetro','Valor'],head)
params[12]=('Penalización por debajo del rango: puntos por unidad',500)
params.append(('Penalización por encima del rango: puntos por unidad',50/.15))
for i,(name,value) in enumerate(params,1):p.write(i,0,name);p.write(i,1,value,percent if i<=12 else None)
p.set_column(0,0,65);p.set_column(1,1,20)
data=[]
d=w.add_worksheet('Datos')
d.write_row(0,0,['Carrera','Crecimiento matriculados','Crecimiento titulados','Búsquedas','Nombre','Profesionales','Vacantes','Alineación','Salario','Competidores','Brecha precio/media','Ventas naturales','Ventas sociedades'],head)
for c in range(4,16):
    a=[source['Hoja1'].cell(1,c).value]+[source['Hoja1'].cell(r,c).value for r in [2,3,4,8,9,10,11,12,13,14,15,16]]
    data.append(a);d.write_row(len(data),0,a)
d.set_column(0,0,40);d.set_column(1,12,22);d.freeze_panes(1,1)
s=w.add_worksheet('Puntajes')
s.write_row(0,0,['Carrera','Matriculados','Titulados','Búsquedas','Nombre','Vacantes/profesionales','Alineación','Salario','Competencia','Oportunidad','Precio','Sector','Demanda','Empleabilidad'],head)
r=w.add_worksheet('Resultados')
r.write_row(0,0,['Carrera','Demanda','Empleabilidad','Oportunidad','Precio','Sector','Índice total','Brecha en puntos vs Negocios','Lectura provisional','Cuadrante de mercado'],head)
q="'Parámetros'!"
base=data[0]
clamp=lambda v:max(0,min(100,v))
relative=lambda x,b:clamp(100*x/b)
rows=[]
for n,a in enumerate(data,2):
    vals=[relative(1+a[1],1+base[1]),relative(1+a[2],1+base[2]),relative(a[3],base[3]),relative(a[4],base[4]),relative(a[6]/a[5],base[6]/base[5]),relative(a[7],base[7]),relative(a[8],base[8]),relative(base[9],a[9])]
    comp=relative(10,a[9])
    if vals[2]>=75:
        opportunity=75+(vals[2]-75)*comp/100 if a[9]<=10 else 50+24*vals[2]/100*comp/100
        quadrant='Mercado atractivo' if a[9]<=10 else 'Mercado competido'
    else:
        opportunity=25+24*vals[2]/75*comp/100 if a[9]<=10 else 24*vals[2]/75*comp/100
        quadrant='Baja tracción' if a[9]<=10 else 'Mercado débil'
    vals += [opportunity,clamp(100-500*max(0,.85-(1+a[10]))-(50/.15)*max(0,(1+a[10])-1.10)),relative(a[11]+a[12],base[11]+base[12])]
    vals += [.65*(.9*vals[2]+.1*vals[3])+.35*(.7*vals[0]+.3*vals[1]),.4*vals[4]+.35*vals[5]+.25*vals[6]]
    formulas=[f'MAX(0,MIN(100,100*(1+Datos!B{n})/(1+Datos!$B$2)))',f'MAX(0,MIN(100,100*(1+Datos!C{n})/(1+Datos!$C$2)))']
    for c in ['D','E']:formulas.append(f'MAX(0,MIN(100,100*Datos!{c}{n}/Datos!${c}$2))')
    formulas += [f'MAX(0,MIN(100,100*(Datos!G{n}/Datos!F{n})/(Datos!$G$2/Datos!$F$2)))']
    for c in ['H','I']:formulas.append(f'MAX(0,MIN(100,100*Datos!{c}{n}/Datos!${c}$2))')
    competition=f'MIN(100,100*{q}$B$19/Datos!J{n})'
    opportunity_formula=f'IF(D{n}>={q}$B$18,IF(Datos!J{n}<={q}$B$19,75+25*(D{n}-{q}$B$18)/(100-{q}$B$18)*({competition})/100,50+24*D{n}/100*({competition})/100),IF(Datos!J{n}<={q}$B$19,25+24*D{n}/{q}$B$18*({competition})/100,24*D{n}/{q}$B$18*({competition})/100))'
    formulas += [f'MAX(0,MIN(100,100*Datos!$J$2/Datos!J{n}))',opportunity_formula,f'MAX(0,MIN(100,100-{q}$B$14*MAX(0,{q}$B$12-(1+Datos!K{n}))-{q}$B$20*MAX(0,(1+Datos!K{n})-{q}$B$13)))',f'MAX(0,MIN(100,100*(Datos!L{n}+Datos!M{n})/(Datos!$L$2+Datos!$M$2)))',f'{q}$B$7*({q}$B$8*D{n}+(1-{q}$B$8)*E{n})+(1-{q}$B$7)*({q}$B$9*B{n}+(1-{q}$B$9)*C{n})',f'{q}$B$10*F{n}+{q}$B$11*G{n}+(1-{q}$B$10-{q}$B$11)*H{n}']
    s.write(n-1,0,a[0])
    for j,(f,v) in enumerate(zip(formulas,vals),1):s.write_formula(n-1,j,'='+f,number,v)
    blocks=[vals[11],vals[12],vals[8],vals[9],vals[10]]
    total=sum(v*wt for v,wt in zip(blocks,[.30,.22,.28,.12,.08]))
    label='Cumple referencia' if total>=90 else 'Brecha moderada' if total>=75 else 'Brecha relevante' if total>=60 else 'Brecha alta'
    r.write(n-1,0,a[0])
    for j,(c,v) in enumerate(zip(['M','N','J','K','L'],blocks),1):r.write_formula(n-1,j,f'=Puntajes!{c}{n}',number,v)
    r.write_formula(n-1,6,f'=B{n}*{q}$B$2+C{n}*{q}$B$3+D{n}*{q}$B$4+E{n}*{q}$B$5+F{n}*{q}$B$6',number,total)
    r.write_formula(n-1,7,f'=G{n}-$G$2',None,total-100)
    r.write_formula(n-1,8,f'=IF(G{n}>={q}$B$15,"Cumple referencia",IF(G{n}>={q}$B$16,"Brecha moderada",IF(G{n}>={q}$B$17,"Brecha relevante","Brecha alta")))',None,label)
    r.write_formula(n-1,9,f'=IF(Puntajes!D{n}>={q}$B$18,IF(Datos!J{n}<={q}$B$19,"Mercado atractivo","Mercado competido"),IF(Datos!J{n}<={q}$B$19,"Baja tracción","Mercado débil"))',None,quadrant)
    rows.append((a[0],total))
for sh in [s,r]:
    sh.set_column(0,0,40);sh.set_column(1,13,21);sh.freeze_panes(1,1);sh.set_row(0,45)
r.set_column(8,9,28);r.autofilter(0,0,12,9)
r.conditional_format('B2:G13',{'type':'3_color_scale','min_type':'num','min_value':0,'mid_type':'num','mid_value':75,'max_type':'num','max_value':100})
g=w.add_worksheet('Guía')
notes=[
('Cambio de versión','Esta versión sustituye la normalización de 70. Negocios=100 por indicador; todo indicador y total se limita a 0–100. Datos comparables según confirmación del usuario. No representa probabilidad de éxito.'),
('Regla de tres','Indicadores favorables: MIN(100,valor/valor_Negocios*100). Competidores: MIN(100,competidores_Negocios/competidores*100). Vacantes/profesionales se compara como una relación.'),
('Crecimientos negativos','Se comparan factores 1+tasa, no tasas directamente: MIN(100,MAX(0,(1+tasa)/(1+tasa_Negocios)*100)). Así una caída obtiene menos puntos. Es una excepción necesaria a la regla de tres de niveles. Tasas base deben ser mayores que -100%.'),
('Oportunidad de mercado','Cuadrantes: atractivo si búsquedas >=75 y competidores <=10; competido si búsquedas >=75 y competidores >10; baja tracción si búsquedas <75 y competidores <=10; débil si búsquedas <75 y competidores >10. El 75 es un umbral, no un peso. Límites incluidos en verde.'),
('Puntuación por cuadrante','Atractivo:75–100; competido:50–74; baja tracción:25–49; débil:0–24. Dentro del cuadrante se gradúa por búsquedas y competencia; ver fórmula en Puntajes. Se prioriza por criterio demanda alta frente a baja tracción. Reglas iniciales editables, no calibradas. Hay saltos al cruzar cuadrantes.'),
('Efecto del tope','Superar a Negocios no agrega puntos por encima de 100. Demanda muy alta no compensa mayor competencia, ni menos competidores compensa demanda baja. Dos carreras por encima de la referencia pueden empatar.'),
('Precio','Precio/media =1+brecha de Hoja1. Rango aceptable 0.85 a 1.10 inclusive:100 puntos. Por debajo se restan 5 puntos por punto porcentual; por encima se restan 3.333333 puntos por punto porcentual. Prima de25% ->50; descuento de25% ->50. Prima de20% ->66.67. Mínimo0. Pendientes editables. No mide capacidad de pago ni margen.'),
('Oportunidad vs demanda','Búsquedas participa en demanda (17.55% total) y en oportunidad (28%). Uso intencional para requerir interés junto con poca oferta; revisar sensibilidad antes de aprobar pesos definitivos.'),
('Pesos','Demanda 30%, empleabilidad 22%, oportunidad 28%, precio 12%, sector 8%. Demanda: nombre 65% (búsquedas90%, rúbrica10%); educativa35% (matriculados70%, titulados30%). Empleabilidad: vacantes/profesionales40%, alineación35%, salario25%.'),
('Sector','Se suman ventas naturales y sociedades; regla de tres frente a Negocios, con tope100. Mide tamaño de actividad asociada, no crecimiento ni rentabilidad del programa.'),
('Fuente de precio','Se conserva Hoja1 fila14 como brecha por carrera. Existen diferencias con PRECIO; no se sustituyeron por promedios de facultad. Las hojas originales quedan copiadas para trazabilidad.'),
('Umbrales iniciales','90–100 cumple referencia;75–<90 brecha moderada;60–<75 brecha relevante;<60 brecha alta. Etiquetas comparativas provisionales, no criterios definitivos de apertura. Una dimensión débil requiere revisión incluso si el total es favorable.'),
('Modalidad: recordar después','Presencial como supuesto descriptivo. Modalidad, entrega y capacidad institucional quedan fuera del cálculo para revisión posterior.'),
('Uso','Editar Datos o Parámetros; fórmulas se recalculan al abrir. Se incluyen valores calculados para previsualización. Mantener Negocios en fila2. El formato % muestra puntajes de 0 a100, no probabilidades.'),
('Validación pendiente','La comparabilidad de los datos se acepta según el usuario. Los pesos y umbrales siguen siendo de criterio; una muestra de mejores carreras no valida predicción. No sustituye evaluación financiera de apertura.'),
]
g.write_row(0,0,['Tema','Explicación'],head);g.set_column(0,0,30);g.set_column(1,1,115)
for i,row in enumerate(notes,1):g.write_row(i,0,row,wrap);g.set_row(i,75)
assert all(0<=v<=100 for _,v in rows)
assert abs(rows[0][1]-100)<1e-9
w.close()
check=openpyxl.load_workbook(path,data_only=True)
assert check['Resultados']['G2'].value==100
assert len(rows)==12
for gap,expected in [(-.25,50),(-.15,100),(.10,100),(.25,50)]:
    actual=clamp(100-500*max(0,.85-(1+gap))-(50/.15)*max(0,(1+gap)-1.10))
    assert abs(actual-expected)<1e-8
print(path)
for name,value in sorted(rows,key=lambda x:-x[1]):print(name,round(value,2))
