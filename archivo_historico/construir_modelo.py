from pathlib import Path
import math
import openpyxl
import xlsxwriter

ROOT = Path(__file__).parent
SOURCE = ROOT.parent / 'Tabla de valores tendencias.xlsx'
src = openpyxl.load_workbook(SOURCE, data_only=True)
out = ROOT / 'Modelo_comparativo_viabilidad.xlsx'
w = xlsxwriter.Workbook(out)
title = w.add_format({'bold': True, 'bg_color': '#006DB5', 'font_color': 'white', 'text_wrap': True})
num = w.add_format({'num_format': '0.00'})
pct = w.add_format({'num_format': '0.00%'})
wrap = w.add_format({'text_wrap': True, 'valign': 'top'})
for sheet in src:
    s = w.add_worksheet('Fuente_' + sheet.title)
    for r, row in enumerate(sheet.values):
        for c, v in enumerate(row):
            if v is not None:
                s.write(r, c, v, title if r == 0 else None)
    s.freeze_panes(1, 3)
    s.set_column(0, 2, 30)
    s.set_column(3, 19, 19)

notes = [
('Alcance', 'Comparativo exploratorio de 11 carreras contra Negocios Internacionales. No estima probabilidad de éxito ni viabilidad financiera de apertura.'),
('Referencia', 'Negocios Internacionales, elegida por el usuario por desempeño satisfactorio y estable. No se asume que sea el promedio estadístico.'),
('Pesos', 'Se conservan 30%, 22%, 28%, 12%, 8%. Competencia pasa íntegramente a oportunidad de mercado; el precio se mide una sola vez en el bloque de 12%. Pagador/capacidad de pago se excluye por falta de datos.'),
('Modalidad pendiente', 'Recordatorio: revisar modalidad y capacidad de implementación en una fase posterior. Presencial como supuesto descriptivo; no entra al cálculo.'),
('Datos utilizados', 'Hoja1: búsquedas locales (fila 4), sin palabras clave ni promedios regional/mundial. Brecha de precio: fila 14. Ventas: suma de filas 15 y 16.'),
('Precio: discrepancias', 'Hoja1 y PRECIO contienen brechas distintas, especialmente Enfermería, Psicología y Multimedia. Se utiliza Hoja1 sin sustituirla con promedios de facultad. Conciliar antes de decidir.'),
('Google Trends: validación pendiente', 'No está documentado si las carreras comparten escala de consulta o un ancla común. Los índices 0–100 se utilizan provisionalmente como comparables; todos los resultados quedan sujetos a validar esa extracción.'),
('Laboral: validación pendiente', 'Negocios tiene 120 profesionales y 1 vacante frente a universos mucho mayores en otras carreras. Verificar equivalencia de búsquedas. Vacantes/profesionales es una señal relativa, no una tasa de empleabilidad.'),
('Economía', 'La base contiene ventas acumuladas de 36 meses, no tasas de crecimiento. El bloque mide tamaño económico asociado, con transformación logarítmica; no crecimiento. Revisar mapeo CIIU y solapamientos.'),
('Normalización', 'Referencia = 70 puntos por indicador. Variables positivas: MIN(100,140*x/(x+referencia)). Crecimientos: limitar a 0–100 el valor 70+50*(x-referencia), usando tasas decimales. No dividir crecimientos negativos.'),
('Competencia y oportunidad', 'Puntaje de menor competencia: MIN(100,140*competidores_ref/(competidores+competidores_ref)). Oportunidad=MIN(puntaje_demanda,RAIZ(puntaje_demanda*puntaje_competencia)). La demanda limita el puntaje: pocos competidores no compensan bajo interés.'),
('Precio', 'Brecha = precio institucional/media de mercado - 1, según Hoja1. Puntaje=limitar a 0–100 [70-50*(brecha-brecha_ref)]. Menor prima mejora competitividad relativa; no demuestra capacidad de pago ni rentabilidad.'),
('Ventas', 'Puntaje=limitar a 0–100 [70+15*LN(ventas/ventas_ref)]. Compresión de escala para no dejar dominar magnitudes de sectores grandes. Parámetro de criterio, no estimado estadísticamente.'),
('Umbrales provisionales', '>=75: superior a referencia; 65 a <75: cercano a referencia; 50 a <65: inferior, requiere ajustes; <50: rezago importante. Son categorías comparativas, no autorización de apertura.'),
('Señal de oportunidad', 'Demanda alta: búsquedas >=1.2 veces referencia; competencia baja: <=0.8 veces referencia. Prioridad condicionada a empleabilidad y precio >=50. Si ambos criterios favorables no coinciden, se muestra la posición respecto a referencia.'),
('Doble uso de búsquedas', 'Búsquedas influye directamente en demanda y también en oportunidad (28%). Es intencional para identificar demanda desatendida; revisar sensibilidad y comparabilidad antes de aprobar estos pesos.'),
('Validación posterior', 'Todas las carreras son casos seleccionados por buen desempeño. Esta muestra no permite validar capacidad predictiva ni fijar umbrales definitivos. Añadir casos medios/bajos y resultados de aperturas.'),
('Edición', 'Datos y Parámetros son editables. Puntajes y Resultados tienen fórmulas y valores calculados guardados. El libro se recalcula al abrir. Mantener la fila 2 de Datos como referencia.'),
]
s=w.add_worksheet('Guía')
s.write_row(0,0,['Tema','Criterio / limitación'],title)
for r,row in enumerate(notes,1):
    s.write_row(r,0,row,wrap); s.set_row(r,64)
s.set_column(0,0,29); s.set_column(1,1,115); s.freeze_panes(1,0)

params=[('Demanda e interés',.30),('Empleabilidad y valor',.22),('Oportunidad de mercado',.28),('Precio frente al mercado',.12),('Tamaño del sector económico',.08),('Demanda del nombre dentro de demanda',.65),('Búsquedas dentro del nombre',.90),('Matriculados dentro de demanda educativa',.70),('Vacantes/profesionales dentro de empleabilidad',.40),('Alineación dentro de empleabilidad',.35),('Ancla de referencia',70),('Pendiente crecimientos y precio',50),('Sensibilidad logarítmica ventas',15),('Umbral superior',75),('Umbral cercano',65),('Umbral rezago',50),('Demanda alta / referencia',1.2),('Competencia baja / referencia',.8),('Mínimo empleabilidad y precio para priorizar',50)]
p=w.add_worksheet('Parámetros'); p.write_row(0,0,['Parámetro','Valor'],title)
for r,(label,value) in enumerate(params,1): p.write(r,0,label); p.write_number(r,1,value,pct if r<=10 else num)
p.write(21,0,'Suma pesos bloques');p.write_formula(21,1,'=SUM(B2:B6)',pct,1)
p.set_column(0,0,58);p.set_column(1,1,18)
d=w.add_worksheet('Datos')
headers=['Carrera','Crecimiento matriculados','Crecimiento titulados','Búsquedas locales','Calificación nombre','Profesionales','Vacantes','Alineación','Salario','Competidores','Brecha precio vs media','Ventas naturales','Ventas sociedades']
d.write_row(0,0,headers,title); d.freeze_panes(1,1);d.set_column(0,0,39);d.set_column(1,12,21);d.set_row(0,45)
data=[]
for col in range(4,16):
    a=[src['Hoja1'].cell(1,col).value]+[src['Hoja1'].cell(r,col).value for r in [2,3,4,8,9,10,11,12,13,14,15,16]]
    data.append(a)
    for c,v in enumerate(a): d.write(len(data),c,v,pct if c in [1,2,7,10] else num if c else None)
d.autofilter(0,0,12,12)
sc=w.add_worksheet('Puntajes'); sh=['Carrera','Matriculados','Titulados','Búsquedas','Nombre','Vacantes/profesionales','Alineación','Salario','Menor competencia','Oportunidad','Precio','Sector','Demanda bloque','Empleabilidad bloque']
sc.write_row(0,0,sh,title); sc.set_column(0,0,39);sc.set_column(1,13,19);sc.freeze_panes(1,1);sc.set_row(0,45)
rws=w.add_worksheet('Resultados');rh=['Carrera','Demanda 30%','Empleabilidad 22%','Oportunidad 28%','Precio 12%','Sector 8%','Índice /100','Diferencia vs referencia','Clasificación provisional','Señal de oportunidad','Estado']
rws.write_row(0,0,rh,title);rws.set_column(0,0,39);rws.set_column(1,7,19);rws.set_column(8,10,45);rws.freeze_panes(1,1);rws.set_row(0,45)
clamp=lambda x:max(0,min(100,x))
pos=lambda x,b:min(100,140*x/(x+b))
b=data[0];results=[]
for i,a in enumerate(data,2):
    scores=[clamp(70+50*(a[1]-b[1])),clamp(70+50*(a[2]-b[2])),pos(a[3],b[3]),pos(a[4],b[4]),pos(a[6]/a[5],b[6]/b[5]),pos(a[7],b[7]),pos(a[8],b[8]),pos(b[9],a[9])]
    scores += [min(scores[2],math.sqrt(scores[2]*scores[7])),clamp(70-50*(a[10]-b[10])),clamp(70+15*math.log((a[11]+a[12])/(b[11]+b[12])))]
    scores += [.65*(.9*scores[2]+.1*scores[3])+.35*(.7*scores[0]+.3*scores[1]),.4*scores[4]+.35*scores[5]+.25*scores[6]]
    q="'Parámetros'!"
    anchor=q+'$B$12'; slope=q+'$B$13'
    formulas=[f'MAX(0,MIN(100,{anchor}+{slope}*(Datos!B{i}-Datos!$B$2)))',f'MAX(0,MIN(100,{anchor}+{slope}*(Datos!C{i}-Datos!$C$2)))']
    for c in ['D','E']:
        formulas.append(f'MIN(100,2*{anchor}*Datos!{c}{i}/(Datos!{c}{i}+Datos!${c}$2))')
    formulas += [f'MIN(100,2*{anchor}*(Datos!G{i}/Datos!F{i})/(Datos!G{i}/Datos!F{i}+Datos!$G$2/Datos!$F$2))']
    for c in ['H','I']:
        formulas.append(f'MIN(100,2*{anchor}*Datos!{c}{i}/(Datos!{c}{i}+Datos!${c}$2))')
    formulas += [f'MIN(100,2*{anchor}*Datos!$J$2/(Datos!J{i}+Datos!$J$2))',f'MIN(D{i},SQRT(D{i}*I{i}))',f'MAX(0,MIN(100,{anchor}-{slope}*(Datos!K{i}-Datos!$K$2)))',f'MAX(0,MIN(100,{anchor}+{q}$B$14*LN((Datos!L{i}+Datos!M{i})/(Datos!$L$2+Datos!$M$2))))',f'{q}$B$7*({q}$B$8*D{i}+(1-{q}$B$8)*E{i})+(1-{q}$B$7)*({q}$B$9*B{i}+(1-{q}$B$9)*C{i})',f'{q}$B$10*F{i}+{q}$B$11*G{i}+(1-{q}$B$10-{q}$B$11)*H{i}']
    sc.write(i-1,0,a[0])
    for c,(f,v) in enumerate(zip(formulas,scores),1):sc.write_formula(i-1,c,'='+f,num,v)
    blocks=[scores[11],scores[12],scores[8],scores[9],scores[10]]
    total=sum(x*y for x,y in zip(blocks,[.30,.22,.28,.12,.08]))
    classification='Superior a referencia' if total>=75 else 'Cercano a referencia' if total>=65 else 'Inferior: requiere ajustes' if total>=50 else 'Rezago importante'
    signal='Oportunidad prioritaria condicionada' if a[3]>=1.2*b[3] and a[9]<=.8*b[9] and min(blocks[1],blocks[3])>=50 else 'Demanda alta / competencia alta' if a[3]>=b[3] and a[9]>b[9] else 'Demanda alta / competencia baja o igual' if a[3]>=b[3] else 'Demanda baja / competencia alta' if a[9]>b[9] else 'Demanda baja / competencia baja o igual'
    rws.write(i-1,0,a[0])
    for c,(col,val) in enumerate(zip(['M','N','J','K','L'],blocks),1):rws.write_formula(i-1,c,f'=Puntajes!{col}{i}',num,val)
    rws.write_formula(i-1,6,f'=B{i}*{q}$B$2+C{i}*{q}$B$3+D{i}*{q}$B$4+E{i}*{q}$B$5+F{i}*{q}$B$6',num,total)
    rws.write_formula(i-1,7,f'=G{i}-$G$2',num,total-70)
    rws.write_formula(i-1,8,f'=IF(G{i}>={q}$B$15,"Superior a referencia",IF(G{i}>={q}$B$16,"Cercano a referencia",IF(G{i}>={q}$B$17,"Inferior: requiere ajustes","Rezago importante")))',None,classification)
    sf=f'IF(AND(Datos!D{i}>=Datos!$D$2*{q}$B$18,Datos!J{i}<=Datos!$J$2*{q}$B$19,C{i}>={q}$B$20,E{i}>={q}$B$20),"Oportunidad prioritaria condicionada",IF(Datos!D{i}>=Datos!$D$2,IF(Datos!J{i}>Datos!$J$2,"Demanda alta / competencia alta","Demanda alta / competencia baja o igual"),IF(Datos!J{i}>Datos!$J$2,"Demanda baja / competencia alta","Demanda baja / competencia baja o igual")))'
    rws.write_formula(i-1,9,'='+sf,None,signal)
    rws.write(i-1,10,'Exploratorio: validar fuentes y escalas')
    results.append((a[0],total,total-70,classification))
rws.autofilter(0,0,12,10)
rws.conditional_format('B2:G13',{'type':'3_color_scale','min_value':0,'mid_value':70,'max_value':100,'min_type':'num','mid_type':'num','max_type':'num'})
w.close()
assert len(results)==12 and abs(results[0][1]-70)<1e-9
check=openpyxl.load_workbook(out,data_only=True)
assert check['Resultados']['G2'].value==70
assert all(0<=row[1]<=100 for row in results)
report=['# Comparativo exploratorio','', 'Referencia: Negocios Internacionales = 70 puntos. Valores sujetos a validar comparabilidad de Trends, universo laboral y brechas de precio. No son probabilidades ni decisiones de apertura.','', '| Carrera | Índice /100 | Diferencia | Lectura |','|---|---:|---:|---|']
for name,total,diff,label in sorted(results,key=lambda x:x[1],reverse=True):report.append(f'| {name} | {total:.2f} | {diff:+.2f} | {label} |')
(ROOT/'Resultados_exploratorios.md').write_text('\n'.join(report),encoding='utf-8')
print(out)
for row in sorted(results,key=lambda x:x[1],reverse=True):print(row[0],round(row[1],2),row[3])
